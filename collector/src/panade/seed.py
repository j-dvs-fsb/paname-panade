"""Import du classeur PP_lieux_v2.xlsx dans une base vide.

Le classeur est le point de départ, pas une source qu'on resynchronise : une
fois la base remplie, c'est elle qui fait foi. L'import est donc refusé sur
une base qui contient déjà des lieux, et il est tout ou rien : toutes les
lignes sont validées avant la première insertion.
"""
from __future__ import annotations

import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from panade import models
from panade.db import est_vide

# Vocabulaire du classeur -> vocabulaire canonique (models.py).
STATUT = {"à vérifier": "a_verifier", "vérifié": "verifie"}
CONFIANCE = {"sûr": "sur", "à confirmer": "a_confirmer"}
BOOLEEN = {"oui": True, "non": False}

# Valeur de image_url qui signifie « utiliser le logo du lieu » (légende de l'onglet).
LOGO_MUSEE = "LOGO_MUSEE"

# En-tête du classeur -> colonne de la base, par onglet.
ONGLETS = {
    "Familles": {
        "modele": models.Famille,
        "cle": "id",
        "legende": True,
        "colonnes": {
            "famille_id": "id", "libellé": "libelle",
            "gratuite_permanent": "gratuite_permanent",
            "gratuite_temporaire": "gratuite_temporaire",
            "age_max": "age_max", "conditions": "conditions",
            "source / à vérifier": "source",
        },
        "ignorees": set(),
    },
    "Lieux": {
        "modele": models.Lieu,
        "cle": "id",
        "legende": False,
        "colonnes": {
            "id": "id", "nom": "nom", "famille": "famille_id",
            "confiance_famille": "confiance_famille",
            "expos_url": "expos_url", "tarifs_url": "tarifs_url",
            "horaires_url": "horaires_url", "logo_url": "logo_url",
            "gratuite_override": "gratuite_override", "statut": "statut",
            "date_verif": "verifie_le", "notes": "notes", "tier": "tier",
            "plateforme_billetterie": "plateforme_billetterie",
            "notes_crawler": "notes_crawler",
        },
        # Formules recalculées depuis l'onglet Familles : rien à importer.
        "ignorees": {"règle permanent (héritée)", "règle temporaire (héritée)", "âge max (hérité)"},
    },
    "Tarifs": {
        "modele": models.Tarif,
        "cle": "lieu_id",
        "legende": True,
        "colonnes": {c: c for c in (
            "lieu_id", "portee", "regime", "audience", "prix", "plein_tarif", "recurrence",
            "justificatif", "source_url", "verifie_le", "verifie_par")},
        "ignorees": set(),
    },
    "Horaires": {
        "modele": models.Horaire,
        "cle": "lieu_id",
        "legende": True,
        "colonnes": {c: c for c in (
            "lieu_id", "jour_num", "jour", "ferme", "ouverture", "fermeture", "nocturne",
            "portee", "derniere_entree_min", "valide_du", "valide_au", "note")},
        "ignorees": set(),
    },
    "Fermetures": {
        "modele": models.Fermeture,
        "cle": "lieu_id",
        "legende": True,
        "colonnes": {c: c for c in ("lieu_id", "date", "libelle", "recurrent_annuel", "note")},
        "ignorees": set(),
    },
    "Expositions": {
        "modele": models.Exposition,
        "cle": "lieu_id",
        "legende": True,
        "colonnes": {
            "lieu_id": "lieu_id", "titre": "titre", "portee": "portee",
            "date_debut": "date_debut", "date_fin": "date_fin",
            "description": "description", "url": "url", "image_url": "image_url",
            "source": "source", "verifie_le": "vu_le",
        },
        # statut : « complet / dates_a_confirmer / image_a_recuperer » n'a pas de colonne,
        # les trous sont visibles dans les champs eux-mêmes. verifie_par : pas de colonne.
        "ignorees": {"statut", "verifie_par"},
    },
}

# Ordre d'insertion (clés étrangères).
ORDRE = ("Familles", "Lieux", "Tarifs", "Horaires", "Fermetures", "Expositions")


class SeedError(Exception):
    """Le classeur est refusé. `erreurs` liste tous les problèmes, avec onglet et ligne."""

    def __init__(self, erreurs: list[str]):
        super().__init__(f"{len(erreurs)} erreur(s) dans le classeur")
        self.erreurs = erreurs


class DejaInitialisee(Exception):
    pass


@dataclass
class Rapport:
    comptes: dict[str, int] = field(default_factory=dict)
    lieux_verifies: int = 0
    legendes_ignorees: int = 0
    logo_du_lieu: int = 0
    avertissements: dict[str, int] = field(default_factory=dict)

    def avertir(self, message: str) -> None:
        self.avertissements[message] = self.avertissements.get(message, 0) + 1


def _cellule(v):
    """Valeur de cellule -> type simple. Une chaîne vide vaut None."""
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _traduire(valeur, table: dict):
    if isinstance(valeur, str):
        return table.get(valeur.casefold(), valeur)
    return valeur


def _nettoyer_url(valeur: str) -> tuple[str | None, str | None]:
    """(url, avertissement). Le classeur mêle parfois du commentaire aux liens."""
    morceaux = valeur.split()
    if re.fullmatch(r"https?://\S+", morceaux[0]):
        if len(morceaux) > 1:
            return morceaux[0], f"annotation retirée du lien : « {' '.join(morceaux[1:])} »"
        return morceaux[0], None
    return None, f"pas une URL, ignoré : « {valeur} »"


def _lire_onglet(wb, nom: str, rapport: Rapport, erreurs: list[str]) -> list[tuple[int, dict]]:
    """Lignes de données d'un onglet : [(numéro de ligne, {colonne de la base: valeur})]."""
    cfg = ONGLETS[nom]
    if nom not in wb.sheetnames:
        erreurs.append(f"onglet « {nom} » absent")
        return []
    lignes = wb[nom].iter_rows(values_only=True)
    entetes = [str(c).strip() if c is not None else "" for c in next(lignes, ())]

    connus = set(cfg["colonnes"]) | cfg["ignorees"]
    manquants = set(cfg["colonnes"]) - set(entetes)
    inconnus = {e for e in entetes if e and e not in connus}
    if manquants or inconnus:
        if manquants:
            erreurs.append(f"{nom} : colonnes absentes : {sorted(manquants)}")
        if inconnus:
            erreurs.append(f"{nom} : colonnes inconnues : {sorted(inconnus)}")
        return []

    sortie = []
    for numero, ligne in enumerate(lignes, start=2):
        brut = {e: _cellule(v) for e, v in zip(entetes, ligne) if e in cfg["colonnes"]}
        if all(v is None for v in brut.values()):
            continue
        cle_sheet = next(e for e, c in cfg["colonnes"].items() if c == cfg["cle"])
        autres = [v for e, v in brut.items() if e != cle_sheet]
        if cfg["legende"] and all(v is None for v in autres):
            # Ligne d'explication sous le tableau : texte dans la première colonne seulement.
            rapport.legendes_ignorees += 1
            continue
        sortie.append((numero, {cfg["colonnes"][e]: v for e, v in brut.items()}))
    return sortie


def _preparer(nom: str, numero: int, ligne: dict, rapport: Rapport) -> dict:
    """Traduit une ligne du classeur dans le vocabulaire canonique."""
    cfg = ONGLETS[nom]
    for champ in list(ligne):
        valeur = ligne[champ]
        if valeur is None:
            continue
        if (champ == "url" or champ.endswith("_url")) and isinstance(valeur, str):
            if nom == "Expositions" and champ == "image_url" and valeur == LOGO_MUSEE:
                ligne[champ] = None
                rapport.logo_du_lieu += 1
                continue
            url, avertissement = _nettoyer_url(valeur)
            ligne[champ] = url
            if avertissement:
                rapport.avertir(f"{nom} {champ} : {avertissement}")
                if url is None and nom == "Lieux" and champ != "logo_url":
                    # Le lien invalide peut contenir un indice utile à `discover`.
                    note = ligne.get("notes_crawler")
                    trace = f"[seed] {champ} invalide dans le classeur : {valeur}"
                    ligne["notes_crawler"] = f"{note} | {trace}" if note else trace
    if nom == "Lieux":
        ligne["statut"] = _traduire(ligne.get("statut"), STATUT)
        ligne["confiance_famille"] = _traduire(ligne.get("confiance_famille"), CONFIANCE)
        if ligne["statut"] is None:
            del ligne["statut"]
    elif nom in ("Horaires", "Fermetures"):
        for champ in ("ferme", "nocturne", "recurrent_annuel"):
            if champ in ligne:
                ligne[champ] = _traduire(ligne[champ], BOOLEEN)
    elif nom == "Expositions":
        # Rien n'est publié sans score de confiance : tout passe d'abord en relecture.
        ligne["etat"] = "en_revue"
    # Les champs vides n'entrent pas : le modèle applique ses valeurs par défaut.
    return {c: v for c, v in ligne.items() if v is not None}


def _formater(erreur: dict) -> str:
    champ = ".".join(str(p) for p in erreur["loc"]) or "(ligne)"
    valeur = erreur.get("input")
    detail = f" (valeur : {valeur!r})" if champ != "(ligne)" and valeur is not None else ""
    return f"{champ} : {erreur['msg']}{detail}"


def lire_classeur(chemin: str | Path) -> tuple[dict[str, list[models.Row]], Rapport]:
    """Lit et valide tout le classeur. Lève SeedError avec toutes les erreurs trouvées."""
    rapport = Rapport()
    erreurs: list[str] = []
    wb = load_workbook(chemin, read_only=True)
    try:
        brutes = {nom: _lire_onglet(wb, nom, rapport, erreurs) for nom in ORDRE}
    finally:
        wb.close()

    valides: dict[str, list[models.Row]] = {nom: [] for nom in ORDRE}
    for nom in ORDRE:
        cfg = ONGLETS[nom]
        for numero, ligne in brutes[nom]:
            try:
                valides[nom].append(cfg["modele"](**_preparer(nom, numero, ligne, rapport)))
            except ValidationError as exc:
                cle = ligne.get(cfg["cle"])
                erreurs += [f"{nom} ligne {numero} ({cle}) : {_formater(e)}" for e in exc.errors()]

    familles = {f.id for f in valides["Familles"]}
    lieux: set[str] = set()
    for lieu in valides["Lieux"]:
        if lieu.id in lieux:
            erreurs.append(f"Lieux : identifiant en double : {lieu.id}")
        lieux.add(lieu.id)
        if lieu.famille_id not in familles:
            erreurs.append(f"Lieux ({lieu.id}) : famille inconnue : {lieu.famille_id}")
    for nom in ("Tarifs", "Horaires", "Fermetures", "Expositions"):
        for ligne in valides[nom]:
            if ligne.lieu_id not in lieux:
                erreurs.append(f"{nom} : lieu inconnu : {ligne.lieu_id}")

    if erreurs:
        raise SeedError(erreurs)
    return valides, rapport


def inserer(con: sqlite3.Connection, valides: dict[str, list[models.Row]]) -> None:
    with con:  # une seule transaction : tout ou rien
        for nom in ORDRE:
            for ligne in valides[nom]:
                valeurs = ligne.values()
                colonnes = ", ".join(f'"{c}"' for c in valeurs)
                marques = ", ".join(f":{c}" for c in valeurs)
                con.execute(f"INSERT INTO {ligne.table} ({colonnes}) VALUES ({marques})", valeurs)


def importer(con: sqlite3.Connection, chemin: str | Path) -> Rapport:
    """Remplit une base vide avec le classeur. Refuse une base qui contient déjà des lieux."""
    if not est_vide(con):
        raise DejaInitialisee("la base contient déjà des lieux")
    valides, rapport = lire_classeur(chemin)
    inserer(con, valides)
    rapport.comptes = {nom: len(valides[nom]) for nom in ORDRE}
    rapport.lieux_verifies = sum(1 for lieu in valides["Lieux"] if lieu.statut == "verifie")
    return rapport
